from django.shortcuts import render, redirect
from catalogo.models import Producto
from inventario.models import Stock
from ventas.models import Venta
from django.db.models import Sum
from .forms import IngresoMercaderiaForm, TrasladoForm
from django.contrib import messages
from .models import Movimiento
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
import openpyxl
from openpyxl.styles import Font, Alignment

@login_required
def dashboard_view(request):
    # 1. Contar cuántos productos distintos existen
    total_productos = Producto.objects.count()
    
    # 2. Sumar todas las cantidades de la tabla Stock
    stock_agregado = Stock.objects.aggregate(total=Sum('cantidad'))
    total_stock = stock_agregado['total'] or 0
    
    # 3. Sumar todo el dinero de las Ventas registradas
    ventas_agregadas = Venta.objects.aggregate(total=Sum('total'))
    total_ventas = ventas_agregadas['total'] or 0

    # 4. Empaquetar los datos para enviarlos al HTML
    context = {
        'total_productos': total_productos,
        'total_stock': total_stock,
        'total_ventas': total_ventas,
    }
    return render(request, 'dashboard.html', context)

# --- NUEVA VISTA PARA EL FORMULARIO DE INGRESO ---
@login_required
def ingreso_mercaderia_view(request):
    if request.method == 'POST':
        form = IngresoMercaderiaForm(request.POST)
        if form.is_valid():
            movimiento = form.save(commit=False)
            movimiento.tipo = 'ENTRADA' # Candado: Forzamos que sea un ingreso
            movimiento.save() # Aquí se dispara la matemática automática hacia la tabla Stock
            return redirect('inventario_visual') # Lo mandamos al catálogo para que vea su stock sumado
    else:
        form = IngresoMercaderiaForm()
    
    return render(request, 'ingreso_mercaderia.html', {'form': form})

@login_required
def historial_movimientos_view(request):
    # Traemos todo el historial, del más nuevo al más viejo
    movimientos = Movimiento.objects.all().order_by('-fecha')
    return render(request, 'historial.html', {'movimientos': movimientos})

@login_required
def realizar_movimiento_view(request):
    if request.method == 'POST':
        form = TrasladoForm(request.POST)
        if form.is_valid():
            try:
                movimiento = form.save(commit=False)
                movimiento.tipo = 'TRASLADO' # Forzamos que sea traslado
                
                # Registramos qué usuario está haciendo la operación
                if request.user.is_authenticated:
                    movimiento.usuario = request.user
                    
                movimiento.save()
                return redirect('historial') # Lo enviamos a ver el kardex si tuvo éxito
                
            except Exception as e:
                # Si falla (ej. por falta de stock), sale alerta roja
                messages.error(request, str(e))
    else:
        form = TrasladoForm()
    
    return render(request, 'movimientos.html', {'form': form})

@login_required
def exportar_kardex_excel(request):
    # 1. Crear el libro de Excel y seleccionar la hoja activa
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kardex de Movimientos"

    # 2. Definir los encabezados de las columnas
    headers = ['ID', 'Fecha', 'Usuario', 'Tipo', 'SKU Producto', 'Nombre Producto', 'Cantidad', 'Ubicación / Ruta', 'Motivo']
    ws.append(headers)

    # Dar formato de negrita y centrado a la primera fila (Encabezados)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')

    # 3. Traer los datos de la base de datos
    movimientos = Movimiento.objects.all().order_by('-fecha')

    # 4. Llenar el Excel fila por fila
    for mov in movimientos:
        # Darle formato a la fecha para Excel
        fecha_str = mov.fecha.strftime("%d/%m/%Y %H:%M")
        usuario_str = mov.usuario.username if mov.usuario else "Sistema"
        
        # Lógica para mostrar la ruta origen/destino según el tipo
        if mov.tipo == 'ENTRADA':
            ruta = f"A: {mov.ubicacion.nombre}"
        elif mov.tipo == 'SALIDA':
            ruta = f"De: {mov.ubicacion.nombre}"
        else:
            ruta = f"De: {mov.ubicacion.nombre} -> A: {mov.ubicacion_destino.nombre if mov.ubicacion_destino else 'N/A'}"

        ws.append([
            mov.id,
            fecha_str,
            usuario_str,
            mov.tipo,
            mov.producto.sku,
            mov.producto.nombre,
            mov.cantidad,
            ruta,
            mov.motivo
        ])

    # 5. Ajustar el ancho de las columnas para que se vea bien (Opcional pero recomendado)
    column_widths = [5, 20, 15, 12, 15, 30, 10, 35, 30]
    for i, column_width in enumerate(column_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = column_width

    # 6. Preparar la respuesta HTTP para que el navegador descargue el archivo
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Kardex_Movimientos.xlsx"'
    wb.save(response)

    return response